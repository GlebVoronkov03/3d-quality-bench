"""Сборка docs/PLER_HQ_presentation.pptx — 10–15 мин, визуал для CV/3D-экспертов."""
from __future__ import annotations

import csv
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "pptx_assets"
OUT = ROOT / "docs" / "PLER_HQ_presentation.pptx"
CAT = ROOT / "DataBase" / "metadata" / "catalog.csv"

W, H = Inches(13.333), Inches(7.5)
NAVY = RGBColor(0x24, 0x30, 0x44)
INK = RGBColor(0x1C, 0x24, 0x30)
ACCENT = RGBColor(0x9C, 0x4A, 0x2B)
MUTED = RGBColor(0x5C, 0x65, 0x70)
CREAM = RGBColor(0xF4, 0xF1, 0xEA)
WHITE = RGBColor(0xFF, 0xFC, 0xF7)
TEAL = RGBColor(0x3D, 0x6B, 0x6B)
LINE = RGBColor(0xD7, 0xD0, 0xC4)

CATEGORIES = [
    ("sampling", "Контроль дискретизации", "6", "аналитический контроль ошибки без маски"),
    ("canonical", "Канонические эталоны", "9", "якоря Stanford / AIM@SHAPE / LIRIS"),
    ("acquisition", "Артефакты оцифровки", "8", "sensor residual, отверстия, плотность"),
    ("thin_feature", "Тонкие структуры", "6", "пряди, спицы, стенки; QEM схлопывает"),
    ("sharp_feature", "Жёсткие рёбра", "8", "G0 и hard-surface; Laplacian стирает crease"),
    ("perceptual", "Перцептивное маскирование", "13", "лицо, identity, saliency тела"),
]

PROV_ICON = {"personal": "●", "project": "◇", "literature": "▲", "open": "■"}
PROV_LEGEND = "● личные   ◇ проект   ▲ литература   ■ CC / синтез"


def _set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def _fill(shape, rgb):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb
    shape.line.fill.background()


def blank(prs: Presentation):
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(bg, CREAM)
    sp_tree = sl.shapes._spTree
    sp = bg._element
    sp_tree.remove(sp)
    sp_tree.insert(2, sp)
    return sl


def footer(sl, page: str, extra: bool = False):
    box = sl.shapes.add_textbox(Inches(0.4), Inches(7.12), Inches(10.5), Inches(0.28))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = "PLER-HQ   ·   доп.   ·   50 × 4 × 10 = 2000" if extra else "PLER-HQ   ·   50 × 4 × 10 = 2000"
    _set_run(r, 11, color=ACCENT if extra else MUTED)
    box2 = sl.shapes.add_textbox(Inches(11.6), Inches(7.12), Inches(1.3), Inches(0.28))
    p = box2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = page
    _set_run(r, 11, color=MUTED)


def title_bar(sl, text: str):
    bar = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.92))
    _fill(bar, NAVY)
    box = sl.shapes.add_textbox(Inches(0.45), Inches(0.18), Inches(12.4), Inches(0.62))
    p = box.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = text
    _set_run(r, 28, True, WHITE)


def add_text(sl, x, y, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _set_run(r, size, bold, color)
    return box


def add_img(sl, path: Path, x, y, w, h=None):
    if not path.exists() or path.stat().st_size < 500:
        return None
    kwargs = {"left": Inches(x), "top": Inches(y), "width": Inches(w)}
    if h is not None:
        kwargs["height"] = Inches(h)
    return sl.shapes.add_picture(str(path), **kwargs)


def notes(sl, text: str):
    sl.notes_slide.notes_text_frame.text = text


def load_catalog():
    with CAT.open(encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))


def by_cat(rows, cat):
    return [r for r in rows if r["category"] == cat]


def card(sl, x, y, w, h):
    sh = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    _fill(sh, WHITE)
    sh.line.color.rgb = LINE
    return sh


def build() -> Path:
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    rows = load_catalog()
    n = 0

    n += 1
    sl = blank(prs)
    band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(band, NAVY)
    add_text(sl, 0.6, 1.7, 12, 0.5, "PLER-HQ", 18, False, ACCENT)
    add_text(sl, 0.6, 2.15, 12, 1.4, "Геометрический FR-корпус\nкачества треугольных мешей", 36, True, WHITE)
    add_text(sl, 0.6, 4.7, 12, 0.5, "50 эталонов   ·   4 оператора   ·   10 уровней   ·   2000 стимулов", 20, False, WHITE)
    add_text(sl, 0.6, 6.45, 12, 0.4, "Computer Vision / 3D geometry processing", 14, False, RGBColor(0xB8, 0xC0, 0xCC))
    notes(sl, "Полный факториал processing-артефактов. Не compression-QA и не textured-mesh QA.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Три оси, которые нельзя смешивать")
    triples = [
        (0.5, "Texture / codec", "Nehmé TOG'23\nSJTU-TMQA · TSMD", "геометрия смешана\nс QP / атласом"),
        (4.7, "Point cloud", "BASICS · SJTU-PCQA\nLS-PCQA", "нет связности\nдругая дискретизация"),
        (8.9, "Processing", "LIRIS · CMDM\nIEETA", "эталонов ≤ 5\nступеней ≤ 4–5"),
    ]
    for x, h1, h2, h3 in triples:
        card(sl, x, 1.3, 3.85, 5.2)
        add_text(sl, x + 0.25, 1.55, 3.4, 0.7, h1, 18, True, NAVY)
        add_text(sl, x + 0.25, 2.45, 3.4, 1.5, h2, 16, False, MUTED)
        add_text(sl, x + 0.25, 4.4, 3.4, 1.6, h3, 16, False, ACCENT)
    footer(sl, str(n), extra=True)
    notes(sl, "SOTA 2025 учится на textured/colored базах. PLER считает геометрию. Point cloud — другая задача.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Покрытие свойств")
    add_img(sl, ASSETS / "fig_coverage.png", 0.25, 1.05, 12.8)
    footer(sl, str(n))
    notes(
        sl,
        "Точка — полное покрытие. Полукруг — частично. Пусто — нет. "
        "PLER-HQ — единственная полная нижняя строка: FR mesh, 4 processing-оператора, "
        "примитивы, сканы, L≥10, N≥50, geom-only.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Существующие корпуса")
    add_img(sl, ASSETS / "fig_compare.png", 0.2, 1.05, 12.9)
    footer(sl, str(n))
    notes(
        sl,
        "+ LIRIS: чистый masking. − 4 рефа, 1 тип. "
        "+ CMDM: DSIS, цвет. − L=4, 5 рефов. "
        "+ Nehmé: масштаб, CS-MOS. − compression confound. "
        "+ TSMD: 42 рефа, codec realism. − только AoM. "
        "PLER-HQ: N=50 как textured-базы, но без texture-confound.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Эталоны")
    add_img(sl, ASSETS / "fig_nref.png", 0.35, 1.35, 12.6, 5.3)
    footer(sl, str(n), extra=True)
    notes(sl, "N=50 в том же порядке, что TSMD 42 и Nehmé 55, без texture-confound.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "+  /  −")
    add_img(sl, ASSETS / "fig_proscons.png", 0.15, 1.1, 13.0)
    footer(sl, str(n), extra=True)
    notes(
        sl,
        "Честный минус PLER-HQ: нет текстуры и codec. Это не пробел спецификации, а граница оси.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Что закрывает PLER-HQ")
    items = [
        "FR-пары  (M_ref, T(M_ref))",
        "4 processing-оператора",
        "L = 10 на тип",
        "6 функциональных классов",
        "сканы + примитивы + CAD",
        "MOS ⊂ факториала",
    ]
    for i, t in enumerate(items):
        col, row = i % 3, i // 3
        x, y = 0.55 + col * 4.2, 1.5 + row * 2.5
        card(sl, x, y, 3.9, 2.15)
        add_text(sl, x + 0.2, y + 0.7, 3.5, 0.9, t, 18, True, NAVY, PP_ALIGN.CENTER)
    footer(sl, str(n))
    notes(sl, "Не заменяет Nehmé по текстуре. Ортогональная ось: geometry processing.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Не покрываем")
    add_img(sl, ASSETS / "fig_notthis.png", 0.15, 1.35, 13.0)
    footer(sl, str(n), extra=True)
    notes(sl, "Граница применимости. HybridMQA 2025 всё ещё учится на textured/colored базах.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Устройство")
    add_img(sl, ASSETS / "fig_architecture.png", 0.15, 1.4, 13.0)
    footer(sl, str(n))
    notes(sl, "2000 мешей — для FR-метрик. MOS берёт подмножество пар. θ параметризует T.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "FR-пара")
    add_img(sl, ASSETS / "fig_principle.png", 0.15, 1.55, 13.0)
    footer(sl, str(n), extra=True)
    notes(sl, "Якорь ref=self. Согласование с MOS: PLCC / SROCC. Не NR, не pairwise T23D.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Согласование с MOS")
    add_img(sl, ASSETS / "fig_metrics.png", 0.15, 1.25, 13.0)
    footer(sl, str(n), extra=True)
    notes(sl, "q считается на полном факториале. s — MOS подмножества. Без RMSE на этом слайде: две ранговые оси.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Таксономия")
    add_img(sl, ASSETS / "fig_taxonomy.png", 0.2, 1.15, 12.9)
    footer(sl, str(n))
    notes(sl, "Квоты сдвинуты под функциональный провал метрики. Сумма = 50. Пунктир — равномерная доля 50/6.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Источники")
    add_img(sl, ASSETS / "fig_licenses.png", 0.2, 1.2, 12.9)
    footer(sl, str(n), extra=True)
    notes(sl, "личные 10, проект 6, литература 24, CC/синтез 10. Сумма 50.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Зачем класс")
    add_img(sl, ASSETS / "fig_roles.png", 0.15, 1.05, 13.0)
    footer(sl, str(n))
    notes(
        sl,
        "Дискретизация — нулевой bias формы. Канон — литература. "
        "Оцифровка — отверстия и sensor. Тонкие — QEM схлопывает пряди. "
        "Рёбра — Laplacian стирает G0. Маска — identity лица и тела.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Почему эти квоты")
    for i, (_k, title, num, why) in enumerate(CATEGORIES):
        col, row = i % 3, i // 3
        x, y = 0.45 + col * 4.2, 1.2 + row * 2.75
        card(sl, x, y, 4.0, 2.55)
        add_text(sl, x + 0.1, y + 0.2, 3.8, 0.5, num, 28, True, ACCENT, PP_ALIGN.CENTER)
        add_text(sl, x + 0.1, y + 0.75, 3.8, 0.55, title, 14, True, NAVY, PP_ALIGN.CENTER)
        add_text(sl, x + 0.15, y + 1.4, 3.7, 0.95, why, 13, False, MUTED, PP_ALIGN.CENTER)
    footer(sl, str(n))
    notes(
        sl,
        "13 на маскирование: лицо и тело — главный провал FR-метрик. "
        "9 канонических — сопоставимость. 8 оцифровки и 8 рёбер — два ортогональных отказа. "
        "По 6 на дискретизацию и тонкие структуры — якорь и stress-тест LOD.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Почему 50 · 4 · 10")
    add_img(sl, ASSETS / "fig_scale.png", 0.15, 1.25, 13.0)
    footer(sl, str(n))
    notes(
        sl,
        "50: порядок textured-баз, 6 функциональных классов. 4: шум, Laplacian, LOD/QEM, композиция. "
        "10: психометрика; CMDM 4 ступени грубо; MOS возьмёт подмножество.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Грани")
    add_img(sl, ASSETS / "fig_span.png", 0.2, 1.2, 12.9)
    footer(sl, str(n), extra=True)
    notes(sl, "От 12 граней куба до ~6e6 у рельефа. Пунктир 10^6 — зона, где FR-метрики дорогие.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Факториал")
    add_img(sl, ASSETS / "fig_factorial.png", 0.25, 1.2, 12.8)
    footer(sl, str(n), extra=True)
    notes(sl, "По 500 стимулов на оператор. Полный декартово произведение, без дыр.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Операторы T_θ")
    add_img(sl, ASSETS / "fig_operators.png", 0.15, 1.2, 13.0)
    footer(sl, str(n))
    notes(
        sl,
        "Шум: sigma_rel * диагональ AABB. Сглаж.: k итераций Laplacian. "
        "LOD: Open3D QEM, доля граней r (прокси геометрии кодека). Гибрид: сначала r, затем sigma той же ступени.",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Гибрид")
    add_img(sl, ASSETS / "fig_hybrid.png", 0.15, 1.45, 13.0)
    footer(sl, str(n), extra=True)
    notes(sl, "Сначала QEM, затем шум того же l. Не полное декартово (r × sigma) — иначе 50×10×10.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Лестница θ_l")
    add_img(sl, ASSETS / "fig_ladder.png", 0.15, 1.15, 13.0)
    footer(sl, str(n), extra=True)
    notes(
        sl,
        "Логарифмическая по sigma и r. Smooth — линейная по k. "
        "Гибрид не вводит новый параметр: пара (r_l, sigma_l).",
    )

    n += 1
    sl = blank(prs)
    title_bar(sl, "Stanford Bunny  ·  читаемый уровень")
    labels = [
        ("ref", "эталон"),
        ("noise", "Шум   σ=0.004"),
        ("smoothing", "Сглаж.   k=5"),
        ("decimation", "LOD   r=0.15"),
        ("combined", "Гибрид   r∘σ"),
    ]
    for i, (key, lab) in enumerate(labels):
        x = 0.28 + i * 2.6
        add_img(sl, ASSETS / "bunny" / f"{key}.png", x, 1.15, 2.48, 2.48)
        add_text(sl, x, 3.7, 2.48, 0.4, lab, 13, True, NAVY, PP_ALIGN.CENTER)
    add_text(
        sl,
        0.5,
        4.35,
        12.3,
        2.2,
        "метод читается, форма сохранена\nкрайние ступени — в корпусе, не в демо",
        16,
        False,
        MUTED,
        PP_ALIGN.CENTER,
    )
    footer(sl, str(n))
    notes(sl, "Шум и гибрид — ступень 4; сглаживание и LOD — ступень 5. Крайние ступени в корпусе, не на слайде.")

    for cat, title, num, why in CATEGORIES:
        n += 1
        sl = blank(prs)
        title_bar(sl, f"{title}   ·   {num}")
        add_text(sl, 0.45, 0.94, 12.4, 0.28, f"проба: {why}      {PROV_LEGEND}", 11, False, MUTED)
        items = by_cat(rows, cat)
        k = len(items)
        cols = 5 if k >= 8 else (4 if k >= 6 else 3)
        nrows = (k + cols - 1) // cols
        cell_w = 12.4 / cols
        cell_h = 5.35 / max(nrows, 1)
        img_w = min(cell_w - 0.16, cell_h - 0.62)
        for i, rec in enumerate(items):
            rr, cc = divmod(i, cols)
            x = 0.45 + cc * cell_w
            y = 1.28 + rr * cell_h
            png = ASSETS / "refs" / f"{rec['slug']}.png"
            add_img(sl, png, x, y, img_w, img_w)
            icon = PROV_ICON.get(rec.get("provenance", ""), "")
            name = rec["name_ru"].split("(")[0].strip()
            fn = rec.get("function_ru", "")
            add_text(sl, x, y + img_w + 0.01, img_w, 0.26, f"{icon} {name}", 11, True, NAVY, PP_ALIGN.CENTER)
            add_text(sl, x, y + img_w + 0.26, img_w, 0.28, fn, 9, False, MUTED, PP_ALIGN.CENTER)
        footer(sl, str(n))
        notes(sl, why)

    n += 1
    sl = blank(prs)
    title_bar(sl, "Геометрия эталонов")
    add_img(sl, ASSETS / "fig_stats.png", 0.15, 1.05, 13.0)
    footer(sl, str(n), extra=True)
    notes(sl, "н/м — доля non-manifold рёбер. крив. — средний двугранный угол. Медиана по классу.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Сводка по эталонам")
    add_img(sl, ASSETS / "fig_stats_table.png", 0.1, 0.98, 13.1, 6.05)
    footer(sl, str(n), extra=True)
    notes(sl, "род только у watertight. диам. — единицы модели, между эталонами несравнимы. Файлы >180 МБ: только верш./гран./диам.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "MOS  ⊂  факториала")
    mos = [
        ("пары", "ref слева · dist справа"),
        ("шкала", "5-балльная similarity"),
        ("стимул", "textureless turntable"),
        ("N", "студенты, 50–300"),
        ("не", "не DSIS / ACR / 2AFC"),
        ("этика", "комплект Воронова (1)"),
    ]
    for i, (a, b) in enumerate(mos):
        col, row = i % 3, i // 3
        x, y = 0.5 + col * 4.2, 1.45 + row * 2.5
        card(sl, x, y, 3.95, 2.2)
        add_text(sl, x + 0.2, y + 0.45, 3.55, 0.5, a, 22, True, ACCENT, PP_ALIGN.CENTER)
        add_text(sl, x + 0.2, y + 1.15, 3.55, 0.7, b, 15, False, NAVY, PP_ALIGN.CENTER)
    footer(sl, str(n), extra=True)
    notes(sl, "Предрендер turntable 10–12 с. Комплект руководителя Mozhaeva не используется.")

    n += 1
    sl = blank(prs)
    title_bar(sl, "Корпус  ⊃  MOS")
    add_img(sl, ASSETS / "fig_mos_nest.png", 0.15, 1.3, 13.0)
    footer(sl, str(n), extra=True)
    notes(sl, "2000 мешей считаются всегда. MOS — выборка пар, не второй корпус.")

    n += 1
    sl = blank(prs)
    band = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    _fill(band, NAVY)
    add_text(sl, 0.6, 1.55, 12, 0.4, "PLER-HQ", 16, False, ACCENT)
    add_text(sl, 0.6, 2.15, 12, 1.0, "50 × 4 × 10 = 2000", 44, True, WHITE)
    add_text(
        sl,
        0.6,
        3.6,
        12,
        2.0,
        "геометрия без texture-confound\nполный факториал processing-артефактов\nMOS — подмножество, не замена корпуса",
        22,
        False,
        WHITE,
    )
    notes(sl, "Ортогональны Nehmé/TSMD. Нужны для PLER-2.0 FR.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    print(build())
